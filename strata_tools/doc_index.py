"""imprint/doc_index.py — cached file indexer for the assistant's ☁ OneDrive
and local-file access.

Ported from Sentinel Forge's proven lyceum/doc_index.py. Each file's text
is extracted ONCE and cached keyed by path + modification time; later runs
reuse the cache and only re-extract changed files. All extractors are
optional-dependency-guarded; nothing here raises — failures degrade to
empty text. Read-only by design.
"""
from __future__ import annotations

import json
import os
import re

try:
    import docx as _docx            # python-docx
except Exception:
    _docx = None
try:
    import pypdf as _pypdf
except Exception:
    _pypdf = None
try:
    from bs4 import BeautifulSoup as _BS
except Exception:
    _BS = None
try:
    import openpyxl as _openpyxl
except Exception:
    _openpyxl = None

SUPPORTED = (".md", ".txt", ".docx", ".pdf", ".html", ".htm",
             ".xlsx", ".xlsm", ".csv")

# Spreadsheets can be enormous; cap rows per sheet.
_XLSX_MAX_ROWS = 1500

# A user folder tree contains git repos, caches, and vendored binaries
# that must never be indexed.
EXCLUDE_DIRS = {".git", "__pycache__", ".claude", "node_modules", "tts",
                "dist", "build", ".venv", "venv", "site-packages"}


def _extract_xlsx(path: str) -> str:
    """Spreadsheet → readable text: 'Sheet: <name>' header, then one line
    per row with cells joined by ' | '. data_only=True returns the last
    CALCULATED value for formula cells."""
    wb = _openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            lines = [f"Sheet: {ws.title}"]
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= _XLSX_MAX_ROWS:
                    lines.append(f"... (first {_XLSX_MAX_ROWS} rows shown)")
                    break
                cells = [str(c).strip() for c in row
                         if c is not None and str(c).strip()]
                if cells:
                    lines.append(" | ".join(cells))
            if len(lines) > 1:
                parts.append("\n".join(lines))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return "\n\n".join(parts)


def extract_text(path: str) -> str:
    """Best-effort plain text from a file. "" on any failure/unsupported."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".md", ".txt", ".csv"):
            with open(path, encoding="utf-8", errors="replace") as f:
                return f.read()
        if ext in (".xlsx", ".xlsm") and _openpyxl is not None:
            return _extract_xlsx(path)
        if ext == ".docx" and _docx is not None:
            return "\n".join(p.text for p in _docx.Document(path).paragraphs)
        if ext == ".pdf" and _pypdf is not None:
            reader = _pypdf.PdfReader(path)
            return "\n".join((pg.extract_text() or "") for pg in reader.pages)
        if ext in (".html", ".htm"):
            with open(path, encoding="utf-8", errors="replace") as f:
                raw = f.read()
            return _BS(raw, "html.parser").get_text(" ") if _BS \
                else re.sub(r"<[^>]+>", " ", raw)
    except Exception:
        return ""
    return ""


def cache_dir() -> str:
    """Index cache location: STRATA_INDEX_DIR override, else the roomy E:
    offload drive when connected (keeps the small C: SSD clear), else
    %LOCALAPPDATA%."""
    override = os.environ.get("STRATA_INDEX_DIR")
    if override:
        return override
    if os.path.isdir("E:\\"):
        return os.path.join("E:\\", "Strata")
    base = (os.environ.get("LOCALAPPDATA")
            or os.path.expanduser(r"~\AppData\Local"))
    return os.path.join(base, "Strata")


def iter_supported_files(root: str, exclude_dirs=EXCLUDE_DIRS):
    """Yield supported files under ``root``, pruning excluded / dot dirs."""
    for dirpath, dirnames, filenames in os.walk(root or ""):
        dirnames[:] = [d for d in dirnames
                       if d.lower() not in exclude_dirs
                       and not d.startswith(".")]
        for fn in filenames:
            if fn.lower().endswith(SUPPORTED):
                yield os.path.join(dirpath, fn)


def build_index_over(root: str, cache_file: str, max_files: int = 4000,
                     exclude_dirs=EXCLUDE_DIRS):
    """Return [(relative_path, text), ...] for supported files under
    ``root``, via a path+mtime extraction cache persisted to
    ``cache_file``. Safe/defensive: never raises."""
    try:
        with open(cache_file, encoding="utf-8") as f:
            cache = json.load(f)
    except Exception:
        cache = {}

    files = []
    for fp in iter_supported_files(root, exclude_dirs):
        files.append(fp)
        if len(files) >= max_files:
            break

    out, new_cache, changed = [], {}, False
    for fp in files:
        try:
            mtime = os.path.getmtime(fp)
        except OSError:
            continue
        ent = cache.get(fp)
        if ent and ent.get("mtime") == mtime and "text" in ent:
            text = ent["text"]
        else:
            text = extract_text(fp) or ""
            changed = True
        new_cache[fp] = {"mtime": mtime, "text": text}
        if text.strip():
            try:
                label = os.path.relpath(fp, root)
            except ValueError:
                label = os.path.basename(fp)
            out.append((label, text))

    if changed or len(new_cache) != len(cache):
        try:
            os.makedirs(os.path.dirname(cache_file), exist_ok=True)
            with open(cache_file, "w", encoding="utf-8") as f:
                json.dump(new_cache, f)
        except OSError:
            pass
    return out


def onedrive_root() -> str:
    return os.environ.get("OneDrive") or os.path.expanduser(r"~\OneDrive")
